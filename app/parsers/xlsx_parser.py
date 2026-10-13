"""Excel workbook parser with row-level serialization.

Extracts worksheets, row/column values, headers, and sheet names.
Represents rows as structured searchable records.
"""

from __future__ import annotations

from pathlib import Path
from typing import List

from app.indexing.models import DocumentChunk, ParsedDocument


def parse_xlsx(file_path: Path, rel_path: str) -> ParsedDocument:
    """Parse an XLSX file across all worksheets into structured row records."""
    try:
        import openpyxl
    except ImportError:
        return ParsedDocument(
            source_path=rel_path,
            file_type="xlsx",
            error="openpyxl is not installed",
        )

    try:
        wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
    except Exception as exc:
        return ParsedDocument(
            source_path=rel_path,
            file_type="xlsx",
            error=f"Corrupt or unreadable XLSX: {type(exc).__name__}: {exc}",
        )

    chunks: List[DocumentChunk] = []

    try:
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            rows_iter = sheet.iter_rows(values_only=True)

            # Find the header row (first non-empty row)
            headers: List[str] = []
            header_row_idx = 0

            for r_idx, row in enumerate(rows_iter, start=1):
                cell_vals = [str(c).strip() if c is not None else "" for c in row]
                if not any(cell_vals):
                    continue

                if not headers:
                    headers = cell_vals
                    header_row_idx = r_idx
                    continue

                # Process data row
                row_dict = {}
                row_parts = []
                for col_idx, val in enumerate(cell_vals):
                    if not val:
                        continue
                    header = headers[col_idx] if col_idx < len(headers) and headers[col_idx] else f"Col{col_idx+1}"
                    row_dict[header] = val
                    row_parts.append(f"{header}: {val}")

                if not row_parts:
                    continue

                # Both key-value format and bar-separated format for maximum searchability
                raw_text = " | ".join(val for val in cell_vals if val)
                structured_text = f"Sheet: {sheet_name} | {' | '.join(row_parts)}"
                combined_text = f"{structured_text}\nRow data: {raw_text}"

                chunks.append(
                    DocumentChunk(
                        chunk_id=f"{rel_path}#{sheet_name}#r{r_idx}",
                        source_path=rel_path,
                        file_type="xlsx",
                        sheet=sheet_name,
                        row=r_idx,
                        section=sheet_name,
                        text=combined_text,
                        metadata={
                            "file": rel_path,
                            "sheet": sheet_name,
                            "row": r_idx,
                            "columns": row_dict,
                            "raw_text": raw_text,
                        },
                        content_type="table_row",
                    )
                )
    except Exception as exc:
        return ParsedDocument(
            source_path=rel_path,
            file_type="xlsx",
            error=f"Error reading sheets: {type(exc).__name__}: {exc}",
        )
    finally:
        wb.close()

    return ParsedDocument(
        source_path=rel_path,
        file_type="xlsx",
        chunks=chunks,
        metadata={
            "sheets": wb.sheetnames if hasattr(wb, "sheetnames") else [],
            "total_chunks": len(chunks),
        },
    )
